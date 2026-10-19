"""
End-to-end integration test exercising the full user workflow:
Directory creation -> Scanning -> Content extraction -> Proposal generation ->
Atomic Renaming -> Report Generation -> Undo Reversal.
"""

from pathlib import Path
import docx
import openpyxl
import pptx

from src.core.scanner import DirectoryScanner
from src.core.models import ScanConfig, NamingConvention, FileStatus
from src.core.rename_engine import RenameEngine
from src.core.undo_manager import UndoManager
from src.storage.history import HistoryManager
from src.reports.report_manager import ReportManager


def test_full_workflow_end_to_end(tmp_path):
    # 1. Create realistic dummy files matching user scenario
    # Word doc
    doc_path = tmp_path / "document_8392.docx"
    doc = docx.Document()
    doc.add_heading("CI CD Deployment Architecture", level=1)
    doc.add_paragraph("Deployment pipelines and staging configurations.")
    doc.save(str(doc_path))

    # Excel sheet
    xlsx_path = tmp_path / "file_final_v2.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quarterly_Financial_Report"
    ws.append(["Department", "Expense", "Quarter"])
    wb.save(str(xlsx_path))

    # Markdown file
    md_path = tmp_path / "notes_random.md"
    md_path.write_text("# Production Readiness Review\n\nAll checklist passed.", encoding="utf-8")

    # Python module
    py_path = tmp_path / "source123.py"
    py_path.write_text('"""API Authentication Module"""\n\nclass AuthServer:\n    pass\n', encoding="utf-8")

    # 2. Configure scan
    config = ScanConfig(root_path=tmp_path, naming_convention=NamingConvention.TITLE_CASE)
    scanner = DirectoryScanner(config)

    # 3. Execute Scan
    items, summary = scanner.scan()
    assert summary.total_files == 4
    assert summary.ready_to_rename == 4

    proposals = {item.original_name: item.new_name for item in items}
    assert proposals["document_8392.docx"] == "Ci_Cd_Deployment_Architecture.docx"
    assert proposals["file_final_v2.xlsx"] == "Quarterly_Financial_Report.xlsx"
    assert proposals["notes_random.md"] == "Production_Readiness_Review.md"
    assert proposals["source123.py"] == "Api_Authentication_Module.py"

    # 4. Execute Renames
    history_mgr = HistoryManager(tmp_path / "history.sqlite3")
    engine = RenameEngine(history_mgr)
    session_id = "test_sess_integration"
    records = engine.execute_renames(items, session_id=session_id)

    assert len(records) == 4
    assert all(r.status == "RENAMED" for r in records)

    # Verify physical file names on disk
    assert not doc_path.exists()
    assert not xlsx_path.exists()
    assert not md_path.exists()
    assert not py_path.exists()

    renamed_doc = tmp_path / "Ci_Cd_Deployment_Architecture.docx"
    renamed_xlsx = tmp_path / "Quarterly_Financial_Report.xlsx"
    renamed_md = tmp_path / "Production_Readiness_Review.md"
    renamed_py = tmp_path / "Api_Authentication_Module.py"

    assert renamed_doc.exists()
    assert renamed_xlsx.exists()
    assert renamed_md.exists()
    assert renamed_py.exists()

    # 5. Generate Reports
    report_mgr = ReportManager(export_dir=tmp_path / "reports")
    csv_p, txt_p, json_p = report_mgr.generate_all_reports(items, config, summary, session_id=session_id)

    assert csv_p.exists()
    assert txt_p.exists()
    assert json_p.exists()

    # 6. Revert via Undo
    undo_mgr = UndoManager(history_mgr)
    undo_result = undo_mgr.undo_session(session_id)

    assert undo_result.reverted_items == 4
    assert undo_result.failed_items == 0

    # Verify original files are back
    assert doc_path.exists()
    assert xlsx_path.exists()
    assert md_path.exists()
    assert py_path.exists()

    assert not renamed_doc.exists()
    assert not renamed_xlsx.exists()
    assert not renamed_md.exists()
    assert not renamed_py.exists()
