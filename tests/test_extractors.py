"""
Tests for specialized content extractors across supported formats.
"""

from pathlib import Path
import json
import zipfile
import docx
import openpyxl
import pptx
from PIL import Image

from src.core.content_extractors import (
    TextExtractor,
    SourceCodeExtractor,
    DocxExtractor,
    XlsxExtractor,
    PptxExtractor,
    ImageExtractor,
    ExecutableExtractor,
    MediaExtractor,
    FontExtractor,
    ArchiveExtractor,
    UniversalFallbackExtractor,
    FallbackExtractor,
)


class TestExtractors:

    def test_markdown_extractor(self, tmp_path):
        md_file = tmp_path / "notes.md"
        md_file.write_text("# Project Architecture Overview\n\nThis system handles distributed tasks.", encoding="utf-8")

        extractor = TextExtractor()
        assert extractor.can_handle(".md")
        res = extractor.extract(md_file)

        assert res.title == "Project Architecture Overview"
        assert res.file_type == "Markdown Document"

    def test_json_extractor(self, tmp_path):
        json_file = tmp_path / "data.json"
        json_file.write_text(json.dumps({"title": "Annual Performance Review", "version": 2}), encoding="utf-8")

        extractor = TextExtractor()
        res = extractor.extract(json_file)

        assert res.title == "Annual Performance Review"
        assert res.file_type == "JSON Data"

    def test_csv_extractor(self, tmp_path):
        csv_file = tmp_path / "export.csv"
        csv_file.write_text("Invoice Number,Client Name,Total Amount,Payment Date\n101,Acme,500,2026-01-01", encoding="utf-8")

        extractor = TextExtractor()
        res = extractor.extract(csv_file)

        assert "Invoice Number Client Name" in res.title
        assert res.file_type == "CSV Spreadsheet"

    def test_python_source_extractor(self, tmp_path):
        py_file = tmp_path / "source123.py"
        py_file.write_text('"""API Authentication Module"""\n\nclass UserAuthenticator:\n    pass\n', encoding="utf-8")

        extractor = SourceCodeExtractor()
        assert extractor.can_handle(".py")
        res = extractor.extract(py_file)

        assert res.title == "API Authentication Module"

    def test_docx_extractor(self, tmp_path):
        doc_path = tmp_path / "file_final_v2.docx"
        doc = docx.Document()
        doc.add_heading("Professional Internship Report", level=1)
        doc.add_paragraph("This report outlines the internship responsibilities and deliverables.")
        doc.save(str(doc_path))

        extractor = DocxExtractor()
        assert extractor.can_handle(".docx")
        res = extractor.extract(doc_path)

        assert res.title == "Professional Internship Report"
        assert res.file_type == "Word Document"

    def test_xlsx_extractor(self, tmp_path):
        xlsx_path = tmp_path / "budget.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Q3_Financial_Summary"
        ws.append(["Category", "Budget", "Actual"])
        wb.save(str(xlsx_path))

        extractor = XlsxExtractor()
        assert extractor.can_handle(".xlsx")
        res = extractor.extract(xlsx_path)

        assert res.title == "Q3_Financial_Summary"
        assert res.file_type == "Excel Spreadsheet"

    def test_pptx_extractor(self, tmp_path):
        pptx_path = tmp_path / "slides.pptx"
        prs = pptx.Presentation()
        slide_layout = prs.slide_layouts[0]  # Title slide
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = "Product Launch Roadmap"
        prs.save(str(pptx_path))

        extractor = PptxExtractor()
        assert extractor.can_handle(".pptx")
        res = extractor.extract(pptx_path)

        assert res.title == "Product Launch Roadmap"
        assert res.file_type == "PowerPoint Presentation"

    def test_image_extractor(self, tmp_path):
        img_path = tmp_path / "photo.png"
        img = Image.new("RGB", (200, 100), color="blue")
        img.save(str(img_path))

        extractor = ImageExtractor()
        assert extractor.can_handle(".png")
        res = extractor.extract(img_path)

        assert "200x100" in res.summary
        assert res.file_type == "PNG Image"

    def test_archive_extractor(self, tmp_path):
        zip_path = tmp_path / "project_bundle.zip"
        with zipfile.ZipFile(str(zip_path), "w") as zf:
            zf.writestr("Deployment_Guide/README.md", "# Guide")
            zf.writestr("Deployment_Guide/config.json", "{}")

        extractor = ArchiveExtractor()
        assert extractor.can_handle(".zip")
        res = extractor.extract(zip_path)

        assert res.title == "Deployment_Guide"
        assert res.file_type == "Archive"

    def test_corrupt_file_handling(self, tmp_path):
        corrupt_doc = tmp_path / "broken.docx"
        corrupt_doc.write_bytes(b"Not a valid zip or docx content")

        extractor = DocxExtractor()
        res = extractor.extract(corrupt_doc)

        assert res.title is None
        assert res.error is not None

    def test_svg_extractor(self, tmp_path):
        svg_file = tmp_path / "illustration.svg"
        svg_file.write_text('<svg xmlns="http://www.w3.org/2000/svg"><title>Vector Graphic Design</title></svg>', encoding="utf-8")

        extractor = TextExtractor()
        assert extractor.can_handle(".svg")
        res = extractor.extract(svg_file)

        assert res.title == "Vector Graphic Design"
        assert res.file_type == "SVG Vector Graphic"
        assert res.is_supported is True

    def test_url_shortcut_extractor(self, tmp_path):
        url_file = tmp_path / "repo.url"
        url_file.write_text("[InternetShortcut]\nURL=https://github.com/microsoft/vscode\n", encoding="utf-8")

        extractor = TextExtractor()
        assert extractor.can_handle(".url")
        res = extractor.extract(url_file)

        assert "Vscode" in res.title or "Github" in res.title
        assert res.is_supported is True

    def test_latex_extractor(self, tmp_path):
        tex_file = tmp_path / "paper.tex"
        tex_file.write_text("\\documentclass{article}\n\\title{Quantum Computing Advancements}\n\\begin{document}\n\\end{document}", encoding="utf-8")

        extractor = TextExtractor()
        assert extractor.can_handle(".tex")
        res = extractor.extract(tex_file)

        assert res.title == "Quantum Computing Advancements"
        assert res.is_supported is True

    def test_vsix_extractor(self, tmp_path):
        vsix_file = tmp_path / "plugin.vsix"
        with zipfile.ZipFile(str(vsix_file), "w") as zf:
            zf.writestr("extension.vsixmanifest", '<PackageManifest><Metadata><DisplayName>Python Debugger</DisplayName></Metadata></PackageManifest>')

        extractor = ExecutableExtractor()
        assert extractor.can_handle(".vsix")
        res = extractor.extract(vsix_file)

        assert res.title == "Python Debugger"
        assert res.is_supported is True

    def test_pe_executable_extractor(self):
        import sys
        py_exe = Path(sys.executable)
        extractor = ExecutableExtractor()
        assert extractor.can_handle(".exe")
        res = extractor.extract(py_exe)

        assert res.is_supported is True
        assert res.title is not None
        assert "Python" in res.title

    def test_wav_audio_extractor(self, tmp_path):
        import wave
        wav_file = tmp_path / "sample.wav"
        with wave.open(str(wav_file), "wb") as wf:
            wf.setnchannels(1)
            wf.setsampwidth(2)
            wf.setframerate(44100)
            wf.writeframes(b"\x00\x00" * 4410)  # 0.1s audio

        extractor = MediaExtractor()
        assert extractor.can_handle(".wav")
        res = extractor.extract(wav_file)

        assert res.is_supported is True
        assert "WAV" in res.file_type
        assert "44100" in res.summary

    def test_mp3_audio_extractor(self, tmp_path):
        mp3_file = tmp_path / "track.mp3"
        # 100 bytes of dummy frame data followed by ID3v1 tag
        title_bytes = b"Midnight Melody" + b"\x00" * 15
        artist_bytes = b"Audio Artist" + b"\x00" * 18
        tag = b"TAG" + title_bytes + artist_bytes + b"\x00" * 65
        mp3_file.write_bytes(b"\xff\xfb\x90\x00" * 25 + tag)

        extractor = MediaExtractor()
        assert extractor.can_handle(".mp3")
        res = extractor.extract(mp3_file)

        assert res.is_supported is True
        assert res.title == "Audio Artist - Midnight Melody"

    def test_font_extractor(self):
        arial_path = Path("C:/Windows/Fonts/arial.ttf")
        extractor = FontExtractor()
        assert extractor.can_handle(".ttf")
        if arial_path.exists():
            res = extractor.extract(arial_path)
            assert res.is_supported is True
            assert res.title == "Arial"
            assert res.file_type == "TrueType Font"

    def test_universal_fallback_extractor(self, tmp_path):
        dat_file = tmp_path / "unknown_system.dat"
        dat_file.write_bytes(b"\xde\xad\xbe\xef" * 50)

        extractor = UniversalFallbackExtractor()
        assert extractor.can_handle(".dat")
        res = extractor.extract(dat_file)

        assert res.is_supported is True
        assert res.file_type == "Data File"
        assert res.title is None

